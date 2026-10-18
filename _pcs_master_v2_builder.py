"""
Build PCS 测试清单 testing checklist - Master v2.xlsx
Adds 6 upstream meter columns (KW A/B/C + KVAR A/B/C) per PCS block.

Final column layout (50 cols):
  A-C  (1-3)  : metadata
  D-I  (4-9)  : command reference (KW A/B/C, KVAR A/B/C)
  J    (10)   : PCS-1  Within Range?
  K-P  (11-16): PCS-1  readings
  Q-V  (17-22): Upstream Meter 1  readings  [NEW]
  W    (23)   : PCS-2  Within Range?
  X-AC (24-29): PCS-2  readings
  AD-AI(30-35): Upstream Meter 2  readings  [NEW]
  AJ   (36)   : PCS-3  Within Range?
  AK-AP(37-42): PCS-3  readings
  AQ-AV(43-48): Upstream Meter 3  readings  [NEW]
  AW   (49)   : Note
  AX   (50)   : Reminder
"""

from pathlib import Path
import openpyxl
from openpyxl.styles import PatternFill, Alignment, Font
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter
import openpyxl.formatting.formatting

FAT_DIR = Path(
    r"C:\Users\JasonOng\OneDrive - Advancer Global Ltd\AST BD"
    r"\HyESys Dept\3. Hardware (PCS.BATT)\PCS"
    r"\sinosoar - current v2.1\FAT - PCS checklist"
)
SRC = FAT_DIR / "PCS 测试清单 testing checklist - Master.xlsx"
DST = FAT_DIR / "PCS 测试清单 testing checklist - Master v2.xlsx"

DATA_START = 14
DATA_END   = 85

# ── Fills (8-char AARRGGBB) ───────────────────────────────────────────────────
AMBER        = PatternFill(patternType="solid", fgColor="FFFFC000")
GREEN_INPUT  = PatternFill(patternType="solid", fgColor="FF92D050")
UM_HDR_FILL  = PatternFill(patternType="solid", fgColor="FF7030A0")
UM_SUB1_FILL = PatternFill(patternType="solid", fgColor="FFB09CC8")
UM_SUB2_FILL = PatternFill(patternType="solid", fgColor="FFDDD0EA")
# CF fills: set bgColor = fgColor for maximum Excel compatibility
GREEN_CF     = PatternFill(patternType="solid", fgColor="FF00B050", bgColor="FF00B050")
RED_CF       = PatternFill(patternType="solid", fgColor="FFFF0000", bgColor="FFFF0000")

WHITE_BOLD = Font(bold=True, color="FFFFFF")
DARK_BOLD  = Font(bold=True)
CENTER     = Alignment(horizontal="center", vertical="center", wrap_text=True)
CENTER_NW  = Alignment(horizontal="center", vertical="center")


def sc(ws, row, col, value=None, fill=None, font=None, align=None):
    c = ws.cell(row=row, column=col)
    if value is not None:
        c.value = value
    if fill is not None:
        c.fill = fill
    if font is not None:
        c.font = font
    if align is not None:
        c.alignment = align
    return c


def safe_merge(ws, r1, c1, r2, c2, value, fill, font, align=CENTER):
    """Merge cells, removing any overlapping existing merges first."""
    to_remove = []
    for mr in list(ws.merged_cells.ranges):
        if (mr.min_row <= r2 and mr.max_row >= r1 and
                mr.min_col <= c2 and mr.max_col >= c1):
            to_remove.append(str(mr))
    for r in to_remove:
        try:
            ws.unmerge_cells(r)
        except Exception:
            pass
    ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)
    cell = ws.cell(row=r1, column=c1)
    if value is not None:
        cell.value = value
    if fill is not None:
        cell.fill = fill
    if font is not None:
        cell.font = font
    if align is not None:
        cell.alignment = align


def add_um_section(ws, meter_idx, col_start):
    """Add all headers for one upstream meter block (rows 4-12 + data fills)."""
    c = col_start

    safe_merge(ws, 4, c, 4, c+5,
               f"上游电表 {meter_idx}\nUpstream Meter {meter_idx}\nInsert MQTT Topic",
               UM_HDR_FILL, WHITE_BOLD)
    for row in range(5, 10):
        safe_merge(ws, row, c, row, c+5, "", UM_HDR_FILL, WHITE_BOLD)

    safe_merge(ws, 10, c, 10, c+5,
               f"上游电表 {meter_idx} Readings\nUpstream Meter {meter_idx} Readings",
               UM_HDR_FILL, WHITE_BOLD)

    safe_merge(ws, 11, c,   11, c+2, "有功KW\nActive Power KW",       UM_SUB1_FILL, DARK_BOLD)
    safe_merge(ws, 11, c+3, 11, c+5, "无功KVAR\nReactive Power KVAR", UM_SUB1_FILL, DARK_BOLD)

    for i, letter in enumerate("ABCABC"):
        sc(ws, 12, c+i, letter, UM_SUB2_FILL, DARK_BOLD, CENTER)

    for row in range(DATA_START, DATA_END + 1):
        for offset in range(6):
            sc(ws, row, c+offset, fill=UM_SUB2_FILL, align=CENTER_NW)


def rebuild_pcs_reading_header(ws, within_col, read_start,
                                row10_text, kw_text, kvar_text):
    """Explicitly rebuild one PCS reading block header rows 10-12.

    Fixes openpyxl insert_cols bug that corrupts multi-column horizontal merges.
    within_col  : column index of the Within Range? cell (single col, merged rows 10-12)
    read_start  : first column of the 6-col reading block
    """
    # Within Range? — single col, vertically merged rows 10-12
    safe_merge(ws, 10, within_col, 12, within_col, "Within\nRange?", AMBER, DARK_BOLD)

    # Row 10: Equipment Reading header — 6 cols merged
    safe_merge(ws, 10, read_start, 10, read_start + 5,
               row10_text, GREEN_INPUT, DARK_BOLD)

    # Row 11: KW (first 3 cols) | KVAR (last 3 cols)
    safe_merge(ws, 11, read_start,     11, read_start + 2, kw_text,   GREEN_INPUT, DARK_BOLD)
    safe_merge(ws, 11, read_start + 3, 11, read_start + 5, kvar_text, GREEN_INPUT, DARK_BOLD)

    # Row 12: A B C  A B C — individual cells
    for i, letter in enumerate("ABCABC"):
        sc(ws, 12, read_start + i, letter, GREEN_INPUT, DARK_BOLD, CENTER)


# ── CF formula builders ───────────────────────────────────────────────────────
def cf_within_range_green(read_start, cmd_start=4, row=DATA_START):
    rs = [get_column_letter(read_start + i) for i in range(6)]
    cs = [get_column_letter(cmd_start  + i) for i in range(6)]
    or_part = ",".join(f"ISNUMBER({r}{row})" for r in rs)
    ifs = ",".join(
        f"IF(ISNUMBER({cs[i]}{row}),AND(ISNUMBER({rs[i]}{row}),"
        f"ABS({cs[i]}{row}-{rs[i]}{row})<=2),TRUE)"
        for i in range(6)
    )
    return f"AND(OR({or_part}),{ifs})"

def cf_within_range_red(read_start, row=DATA_START):
    rs = [get_column_letter(read_start + i) for i in range(6)]
    return f"OR({','.join(f'ISNUMBER({r}{row})' for r in rs)})"

def cf_reading_red(read_col, cmd_col, row=DATA_START):
    r, c = get_column_letter(read_col), get_column_letter(cmd_col)
    return f"AND(ISNUMBER({r}{row}),ISNUMBER({c}{row}),ABS({c}{row}-{r}{row})>2)"

def cf_um_green(read_col, cmd_col, row=DATA_START):
    r, c = get_column_letter(read_col), get_column_letter(cmd_col)
    return f"AND(ISNUMBER({r}{row}),ISNUMBER({c}{row}),ABS({c}{row}-{r}{row})<=2)"

def cf_um_red(read_col, cmd_col, row=DATA_START):
    r, c = get_column_letter(read_col), get_column_letter(cmd_col)
    return f"AND(ISNUMBER({r}{row}),ISNUMBER({c}{row}),ABS({c}{row}-{r}{row})>2)"

def rng(col_a, col_b=None):
    ca = get_column_letter(col_a) if isinstance(col_a, int) else col_a
    if col_b is None:
        return f"{ca}{DATA_START}:{ca}{DATA_END}"
    cb = get_column_letter(col_b) if isinstance(col_b, int) else col_b
    return f"{ca}{DATA_START}:{cb}{DATA_END}"


# ── MAIN ──────────────────────────────────────────────────────────────────────
print(f"Loading: {SRC.name}")
wb = openpyxl.load_workbook(SRC)
ws = wb.active

# ── Step 0: Capture PCS reading header text from original BEFORE inserts ──────
# Original positions (before any column insertions):
#   PCS-1: Within Range? col 10, reading cols 11-16 (K-P)
#   PCS-2: Within Range? col 17, reading cols 18-23 (R-W)
#   PCS-3: Within Range? col 24, reading cols 25-30 (Y-AD)
# Row 10 = Equipment Reading header (merged across 6 reading cols)
# Row 11 = KW sub-header (cols +0..+2) | KVAR sub-header (cols +3..+5)

def _cell_text(ws, row, col, default=""):
    v = ws.cell(row=row, column=col).value
    return v if (v is not None and str(v).strip()) else default

_DEF_HDR  = "设备读数\nEquipment Reading"
_DEF_KW   = "有功KW\nActive Power KW"
_DEF_KVAR = "无功KVAR\nReactive Power KVAR"

pcs1_hdr  = _cell_text(ws, 10, 11, _DEF_HDR)
pcs1_kw   = _cell_text(ws, 11, 11, _DEF_KW)
pcs1_kvar = _cell_text(ws, 11, 14, _DEF_KVAR)

pcs2_hdr  = _cell_text(ws, 10, 18, _DEF_HDR)
pcs2_kw   = _cell_text(ws, 11, 18, _DEF_KW)
pcs2_kvar = _cell_text(ws, 11, 21, _DEF_KVAR)

pcs3_hdr  = _cell_text(ws, 10, 25, _DEF_HDR)
pcs3_kw   = _cell_text(ws, 11, 25, _DEF_KW)
pcs3_kvar = _cell_text(ws, 11, 28, _DEF_KVAR)

print(f"  Captured PCS header text from source")
print(f"    PCS-1 row10: {repr(pcs1_hdr[:30])}")
print(f"    PCS-2 row10: {repr(pcs2_hdr[:30])}")
print(f"    PCS-3 row10: {repr(pcs3_hdr[:30])}")

# ── Step 1: Insert upstream meter columns (right → left for stable indices) ───
ws.insert_cols(31, 6)   # UM-3: after original AD (col 30)
ws.insert_cols(24, 6)   # UM-2: after original W  (col 23)
ws.insert_cols(17, 6)   # UM-1: after original P  (col 16)
print(f"  Inserted 18 columns. Total: {ws.max_column} ({get_column_letter(ws.max_column)})")

# ── Step 2: Column widths for UM columns — set AFTER inserts ─────────────────
for col in [*range(17, 23), *range(30, 36), *range(43, 49)]:
    ws.column_dimensions[get_column_letter(col)].width = 6.5

# ── Step 3: Static fills for PCS data rows 14-85 ─────────────────────────────
for row in range(DATA_START, DATA_END + 1):
    sc(ws, row, 10, fill=AMBER,       align=CENTER_NW)   # J  = PCS-1 Within Range?
    sc(ws, row, 23, fill=AMBER,       align=CENTER_NW)   # W  = PCS-2 Within Range?
    sc(ws, row, 36, fill=AMBER,       align=CENTER_NW)   # AJ = PCS-3 Within Range?
    for offset in range(6):
        sc(ws, row, 11+offset, fill=GREEN_INPUT, align=CENTER_NW)   # K-P  PCS-1
        sc(ws, row, 24+offset, fill=GREEN_INPUT, align=CENTER_NW)   # X-AC PCS-2
        sc(ws, row, 37+offset, fill=GREEN_INPUT, align=CENTER_NW)   # AK-AP PCS-3

# ── Step 4: Add upstream meter sections (headers + data fills) ────────────────
add_um_section(ws, 1, 17)    # Q-V   (17-22)
add_um_section(ws, 2, 30)    # AD-AI (30-35)
add_um_section(ws, 3, 43)    # AQ-AV (43-48)

# ── Step 5: Explicitly rebuild PCS-1/2/3 header rows 10-12 ───────────────────
# insert_cols has a known bug where it does not correctly shift multi-column
# horizontal merges (e.g. R10:W10 → X10:AC10). We rebuild all three sections
# explicitly to guarantee correct merges, text, and fill after inserts.

# PCS-1: Within Range? col J(10), reading cols K-P (11-16)
rebuild_pcs_reading_header(ws, 10, 11, pcs1_hdr, pcs1_kw, pcs1_kvar)

# PCS-2: Within Range? col W(23), reading cols X-AC (24-29)
rebuild_pcs_reading_header(ws, 23, 24, pcs2_hdr, pcs2_kw, pcs2_kvar)

# PCS-3: Within Range? col AJ(36), reading cols AK-AP (37-42)
rebuild_pcs_reading_header(ws, 36, 37, pcs3_hdr, pcs3_kw, pcs3_kvar)

# ── Step 6: Row heights — prevent UM multi-line content from auto-expanding ───
# Rows 10-11 have multi-line header content in the UM sections.
# Without explicit heights openpyxl expands them to 48 and 75 pt.
ws.row_dimensions[10].height = 35
ws.row_dimensions[11].height = 35
# Row 12 keeps its original height (set by the source file).

# ── Step 7: Rebuild all conditional formatting ────────────────────────────────
ws.conditional_formatting = openpyxl.formatting.formatting.ConditionalFormattingList()

# PCS-1 Within Range? J(10) reads K-P (11-16)
ws.conditional_formatting.add(
    rng(10), FormulaRule([cf_within_range_green(11)], fill=GREEN_CF, stopIfTrue=True))
ws.conditional_formatting.add(
    rng(10), FormulaRule([cf_within_range_red(11)],   fill=RED_CF,   stopIfTrue=True))
# PCS-1 readings K-P (11-16): RED when reading deviates >2 from command
ws.conditional_formatting.add(
    rng(11, 16), FormulaRule([cf_reading_red(11, 4)], fill=RED_CF, stopIfTrue=True))
# Upstream Meter 1 Q-V (17-22): GREEN/RED vs command
ws.conditional_formatting.add(
    rng(17, 22), FormulaRule([cf_um_green(17, 4)], fill=GREEN_CF, stopIfTrue=True))
ws.conditional_formatting.add(
    rng(17, 22), FormulaRule([cf_um_red(17, 4)],   fill=RED_CF,   stopIfTrue=True))

# PCS-2 Within Range? W(23) reads X-AC (24-29)
ws.conditional_formatting.add(
    rng(23), FormulaRule([cf_within_range_green(24)], fill=GREEN_CF, stopIfTrue=True))
ws.conditional_formatting.add(
    rng(23), FormulaRule([cf_within_range_red(24)],   fill=RED_CF,   stopIfTrue=True))
# PCS-2 readings X-AC (24-29): RED when out of range
ws.conditional_formatting.add(
    rng(24, 29), FormulaRule([cf_reading_red(24, 4)], fill=RED_CF, stopIfTrue=True))
# Upstream Meter 2 AD-AI (30-35): GREEN/RED vs command
ws.conditional_formatting.add(
    rng(30, 35), FormulaRule([cf_um_green(30, 4)], fill=GREEN_CF, stopIfTrue=True))
ws.conditional_formatting.add(
    rng(30, 35), FormulaRule([cf_um_red(30, 4)],   fill=RED_CF,   stopIfTrue=True))

# PCS-3 Within Range? AJ(36) reads AK-AP (37-42)
ws.conditional_formatting.add(
    rng(36), FormulaRule([cf_within_range_green(37)], fill=GREEN_CF, stopIfTrue=True))
ws.conditional_formatting.add(
    rng(36), FormulaRule([cf_within_range_red(37)],   fill=RED_CF,   stopIfTrue=True))
# PCS-3 readings AK-AP (37-42): RED when out of range
ws.conditional_formatting.add(
    rng(37, 42), FormulaRule([cf_reading_red(37, 4)], fill=RED_CF, stopIfTrue=True))
# Upstream Meter 3 AQ-AV (43-48): GREEN/RED vs command
ws.conditional_formatting.add(
    rng(43, 48), FormulaRule([cf_um_green(43, 4)], fill=GREEN_CF, stopIfTrue=True))
ws.conditional_formatting.add(
    rng(43, 48), FormulaRule([cf_um_red(43, 4)],   fill=RED_CF,   stopIfTrue=True))

# ── Step 8: Save ──────────────────────────────────────────────────────────────
wb.save(DST)
print(f"Saved: {DST.name}")
print()
print("Column layout:")
print("  J(10)=PCS-1 Within Range?   K-P(11-16)=PCS-1    Q-V(17-22)=UM-1")
print("  W(23)=PCS-2 Within Range?   X-AC(24-29)=PCS-2   AD-AI(30-35)=UM-2")
print("  AJ(36)=PCS-3 Within Range?  AK-AP(37-42)=PCS-3  AQ-AV(43-48)=UM-3")
print("  AW(49)=Note  AX(50)=Reminder")
