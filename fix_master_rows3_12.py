"""
fix_master_rows3_12.py
----------------------
Copies formatting (fill, font, alignment, border, row height, merges, values)
from the reference file rows 3-12 to the master file rows 3-12,
accounting for 3 deleted columns in the master (J=10, W=23, AJ=36).

Column mapping:
  Reference cols 1-9   -> Master cols 1-9   (no shift)
  Reference cols 11-22 -> Master cols 10-21 (shift -1)
  Reference cols 24-35 -> Master cols 22-33 (shift -2)
  Reference cols 37+   -> Master cols 34+   (shift -3)

Deleted reference columns (10, 23, 36) are skipped entirely.
"""

import copy
import openpyxl
from openpyxl.cell import MergedCell
from openpyxl.utils import get_column_letter, column_index_from_string
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side, GradientFill

# ── File paths ────────────────────────────────────────────────────────────────
MASTER_PATH = (
    r"C:\Users\JasonOng\OneDrive - Advancer Global Ltd\AST BD\HyESys Dept"
    r"\3. Hardware (PCS.BATT)\PCS\sinosoar - current v2.1"
    r"\FAT - PCS checklist\PCS 测试清单 testing checklist - Master v2.xlsx"
)
REF_PATH = (
    r"C:\Users\JasonOng\OneDrive - Advancer Global Ltd\AST BD\HyESys Dept"
    r"\3. Hardware (PCS.BATT)\PCS\sinosoar - current v2.1"
    r"\FAT - PCS checklist\tested files"
    r"\PCS FAT Multi 20261014BY001+20261014BY002+26022703840003 2026-10-22 21-25.xlsx"
)

ROW_START = 3
ROW_END = 12

# Deleted columns in the REFERENCE (1-based)
DELETED_REF_COLS = {10, 23, 36}   # J, W, AJ

# ── Column mapping helper ─────────────────────────────────────────────────────
def ref_col_to_master_col(ref_col):
    """Return the master column index for a reference column, or None if deleted."""
    if ref_col in DELETED_REF_COLS:
        return None
    shift = sum(1 for d in DELETED_REF_COLS if d < ref_col)
    return ref_col - shift


# ── Style copy helpers ────────────────────────────────────────────────────────
def is_merged_slave(cell):
    """Return True if the cell is a MergedCell (slave, non-top-left)."""
    return isinstance(cell, MergedCell)


def copy_fill(cell_from, cell_to):
    if is_merged_slave(cell_to):
        return
    if cell_from.fill and cell_from.fill.fill_type not in (None, "none"):
        cell_to.fill = copy.copy(cell_from.fill)
    else:
        cell_to.fill = PatternFill()  # clear


def copy_font(cell_from, cell_to):
    if is_merged_slave(cell_to):
        return
    if cell_from.font:
        cell_to.font = copy.copy(cell_from.font)


def copy_alignment(cell_from, cell_to):
    if is_merged_slave(cell_to):
        return
    if cell_from.alignment:
        cell_to.alignment = copy.copy(cell_from.alignment)


def copy_border(cell_from, cell_to):
    if is_merged_slave(cell_to):
        return
    if cell_from.border:
        cell_to.border = copy.copy(cell_from.border)


# ── Load workbooks ────────────────────────────────────────────────────────────
print("Loading workbooks …")
ref_wb = openpyxl.load_workbook(REF_PATH)
master_wb = openpyxl.load_workbook(MASTER_PATH)

ref_ws = ref_wb.active
master_ws = master_wb.active

print(f"Reference sheet: '{ref_ws.title}'  |  max_col={ref_ws.max_column}  max_row={ref_ws.max_row}")
print(f"Master sheet:    '{master_ws.title}'  |  max_col={master_ws.max_column}  max_row={master_ws.max_row}")

# ── STEP 1: Inspect and print reference rows 3-12 ─────────────────────────────
print("\n" + "="*70)
print("REFERENCE FILE — Rows 3-12 inspection")
print("="*70)

# Merged cells in reference rows 3-12
print("\n--- Merged cell ranges in rows 3-12 ---")
ref_merges_in_range = []
for merge in ref_ws.merged_cells.ranges:
    if merge.min_row <= ROW_END and merge.max_row >= ROW_START:
        ref_merges_in_range.append(merge)
        print(f"  {merge}")

print(f"\nTotal merges in rows 3-12: {len(ref_merges_in_range)}")

# Print per-cell details
print("\n--- Per-cell details (rows 3-12) ---")
for r in range(ROW_START, ROW_END + 1):
    for c in range(1, ref_ws.max_column + 1):
        cell = ref_ws.cell(row=r, column=c)
        if is_merged_slave(cell):
            continue  # slave cells carry no independent styles

        val = cell.value
        fill = cell.fill
        font = cell.font
        align = cell.alignment

        fill_info = "none"
        if fill and fill.fill_type not in (None, "none"):
            fg = fill.fgColor.rgb if fill.fgColor else "N/A"
            fill_info = f"type={fill.fill_type} fg={fg}"

        font_info = "none"
        if font:
            font_info = (f"bold={font.bold} size={font.size} "
                         f"color={font.color.rgb if font.color else 'N/A'} "
                         f"name={font.name}")

        align_info = "none"
        if align:
            align_info = (f"h={align.horizontal} v={align.vertical} "
                          f"wrap={align.wrap_text}")

        if val is not None or (fill and fill.fill_type not in (None, "none")):
            print(f"  R{r}C{c} ({get_column_letter(c)}{r}): "
                  f"val={repr(val)!s:30s} fill=[{fill_info}] "
                  f"font=[{font_info}] align=[{align_info}]")

# Row heights
print("\n--- Row heights (rows 3-12) ---")
for r in range(ROW_START, ROW_END + 1):
    rd = ref_ws.row_dimensions[r]
    print(f"  Row {r}: height={rd.height}")

# ── STEP 2: Unmerge existing merged cells in master rows 3-12 ────────────────
print("\n" + "="*70)
print("MASTER FILE — Unmerging rows 3-12")
print("="*70)

merges_to_remove = []
for merge in list(master_ws.merged_cells.ranges):
    if merge.min_row <= ROW_END and merge.max_row >= ROW_START:
        merges_to_remove.append(str(merge))

for m in merges_to_remove:
    master_ws.merged_cells.remove(m)
    print(f"  Unmerged: {m}")

print(f"Removed {len(merges_to_remove)} existing merges from master rows 3-12.")

# ── STEP 3: Copy merges from reference to master ──────────────────────────────
print("\n" + "="*70)
print("Copying merged cell ranges to master …")
print("="*70)

copied_merges = []
skipped_merges = []

for merge in ref_merges_in_range:
    # Check if any col in the merge is a deleted ref col
    cols_in_merge = set(range(merge.min_col, merge.max_col + 1))
    deleted_in_merge = cols_in_merge & DELETED_REF_COLS

    if deleted_in_merge:
        skipped_merges.append((str(merge), f"spans deleted col(s) {deleted_in_merge}"))
        print(f"  SKIPPED: {merge}  — spans deleted col(s) {deleted_in_merge}")
        continue

    m_min_col = ref_col_to_master_col(merge.min_col)
    m_max_col = ref_col_to_master_col(merge.max_col)

    if m_min_col is None or m_max_col is None:
        skipped_merges.append((str(merge), "maps to None"))
        print(f"  SKIPPED: {merge}  — maps to None master col")
        continue

    # Build master merge string
    min_cell = f"{get_column_letter(m_min_col)}{merge.min_row}"
    max_cell = f"{get_column_letter(m_max_col)}{merge.max_row}"
    master_range = f"{min_cell}:{max_cell}"

    master_ws.merge_cells(master_range)
    copied_merges.append((str(merge), master_range))
    print(f"  Copied: {merge}  ->  {master_range}")

print(f"\nMerges copied: {len(copied_merges)}  |  Skipped: {len(skipped_merges)}")

# ── STEP 4: Copy fill, font, alignment, border, values per cell ───────────────
print("\n" + "="*70)
print("Copying cell styles and values …")
print("="*70)

# Preserve master R3C4 (date cell) value
DATE_ROW, DATE_COL = 3, 4
master_date_value = master_ws.cell(row=DATE_ROW, column=DATE_COL).value
print(f"  Preserving master date cell R3C4 value: {repr(master_date_value)}")

cells_styled = 0
for r in range(ROW_START, ROW_END + 1):
    for ref_c in range(1, ref_ws.max_column + 1):
        master_c = ref_col_to_master_col(ref_c)
        if master_c is None:
            continue  # skip deleted columns

        ref_cell = ref_ws.cell(row=r, column=ref_c)
        m_cell = master_ws.cell(row=r, column=master_c)

        # Reference cell may be a MergedCell slave — styles/values not on slaves;
        # only copy from non-slave reference cells (top-left of their merge group)
        if not is_merged_slave(ref_cell):
            copy_fill(ref_cell, m_cell)
            copy_font(ref_cell, m_cell)
            copy_alignment(ref_cell, m_cell)
            copy_border(ref_cell, m_cell)

            # Copy value, but protect master date cell and skip merged slaves
            if is_merged_slave(m_cell):
                pass  # cannot set value on merged slave cells
            elif r == DATE_ROW and master_c == DATE_COL:
                pass  # keep master value
            else:
                m_cell.value = ref_cell.value

        cells_styled += 1

print(f"  Styled {cells_styled} cells.")

# Restore master date
master_ws.cell(row=DATE_ROW, column=DATE_COL).value = master_date_value
print(f"  Restored master date: {repr(master_date_value)}")

# ── STEP 5: Copy row heights ──────────────────────────────────────────────────
print("\n" + "="*70)
print("Copying row heights …")
print("="*70)

for r in range(ROW_START, ROW_END + 1):
    ref_height = ref_ws.row_dimensions[r].height
    master_ws.row_dimensions[r].height = ref_height
    print(f"  Row {r}: height set to {ref_height}")

# ── STEP 6: Save ──────────────────────────────────────────────────────────────
print("\n" + "="*70)
print("Saving master file …")
print("="*70)
master_wb.save(MASTER_PATH)
print(f"  Saved: {MASTER_PATH}")

# ── Summary ───────────────────────────────────────────────────────────────────
print("\n" + "="*70)
print("SUMMARY")
print("="*70)
print(f"  Merges unmerged from master:  {len(merges_to_remove)}")
print(f"  Merges copied to master:      {len(copied_merges)}")
print(f"  Merges skipped (deleted col): {len(skipped_merges)}")
if skipped_merges:
    for ref_r, reason in skipped_merges:
        print(f"    - {ref_r}: {reason}")
print(f"  Cells styled:                 {cells_styled}")
print(f"  Date cell preserved:          R3C4 = {repr(master_date_value)}")
print("\nDone.")
